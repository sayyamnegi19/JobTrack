"""
Test suite for the Resume ATS app.

External services (Gemini API, job-page fetching, DNS) are mocked everywhere
so the suite is fast, free and offline. Run with:

    venv/bin/python manage.py test resumes

Note: the Gemini embedding/model calls are never made for real here; the
tests only pin down OUR logic (validation, normalization, mapping, guards).
Live model quality is covered by the manual smoke test in the README.
"""

import json
import os
import shutil
import tempfile
from io import BytesIO
from types import SimpleNamespace
from unittest import mock

import requests
from django.conf import settings
from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import SimpleTestCase, TestCase, override_settings
from django.urls import reverse
from docx import Document
from google.genai import errors

from applications.models import JobApplication

from . import parsers
from .ai_client import AIAnalysisError
from .analyzer import _truncate, analyze_resume, build_prompt
from .forms import ResumeAnalysisForm
from .models import Resume, ResumeAnalysis
from .scraper import (
    JobFetchError,
    extract_job_description,
    fetch_job_description,
    fetch_job_page,
    validate_public_url,
)
from .similarity import _cosine_similarity, semantic_similarity

User = get_user_model()

TEMP_MEDIA = tempfile.mkdtemp(prefix="jobtrack_test_media_")


def tearDownModule():
    shutil.rmtree(TEMP_MEDIA, ignore_errors=True)


# ---------------------------------------------------------------------------
# Shared helpers
# ---------------------------------------------------------------------------

RESUME_TEXT = "John Doe, Software Engineer. Python, Django, PostgreSQL. " * 10


def build_text_pdf(text):
    """Build a minimal single-page PDF with correct xref offsets (no deps)."""
    escaped = text.replace("\\", r"\\").replace("(", r"\(").replace(")", r"\)")
    stream = f"BT /F1 12 Tf 72 720 Td ({escaped}) Tj ET".encode("latin-1")

    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n"
        + stream + b"\nendstream",
    ]

    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for number, obj in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n".encode() + obj + b"\nendobj\n"

    xref_pos = len(out)
    out += f"xref\n0 {len(objects) + 1}\n".encode()
    out += b"0000000000 65535 f \n"
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    out += (
        b"trailer\n<< /Size " + str(len(objects) + 1).encode() + b" /Root 1 0 R >>\n"
        b"startxref\n" + str(xref_pos).encode() + b"\n%%EOF\n"
    )
    return bytes(out)


def fake_result_json(overall=72):
    """A JSON payload shaped exactly like the model's response schema."""
    return json.dumps(
        {
            "overall_score": overall,
            "summary": "Solid resume with room to grow.",
            "categories": [
                {"key": "keyword_match", "label": "Keyword Match", "score": 65, "feedback": "Add CI/CD."},
                {"key": "formatting", "label": "ATS Formatting", "score": 90, "feedback": "Clean layout."},
                {"key": "skills_sections", "label": "Skills & Sections", "score": 70, "feedback": "Generic skills."},
                {"key": "experience_impact", "label": "Experience Impact", "score": 60, "feedback": "Add numbers."},
                {"key": "readability", "label": "Readability", "score": 80, "feedback": "Concise."},
            ],
            "matched_keywords": ["Python", " python ", "Django"],
            "missing_keywords": ["Kubernetes"],
            "strengths": ["Quantified achievements."],
            "improvements": [{"priority": "high", "suggestion": "Add testing evidence."}],
        }
    )


def make_mock_client(output_text, prompt_tokens=100, output_tokens=50):
    interaction = mock.Mock()
    interaction.output_text = output_text
    interaction.usage = mock.Mock(
        total_input_tokens=prompt_tokens, total_output_tokens=output_tokens
    )
    client = mock.Mock()
    client.interactions.create.return_value = interaction
    return client


def fake_analyze(resume_text, job_description=None):
    """Stand-in for resumes.analyzer.analyze_resume in view tests."""
    return {
        "mode": "JOB_MATCH" if job_description else "GENERAL",
        "overall_score": 55,
        "result": json.loads(fake_result_json(55)),
        "ai_model": "mock-model",
        "prompt_tokens": 10,
        "output_tokens": 20,
        "semantic_similarity": 0.5 if job_description else None,
    }


def make_http_response(
    status=200,
    body=b"",
    headers=None,
    chunks=None,
    is_redirect=False,
    encoding="utf-8",
):
    """A fake requests.Response good enough for the scraper's usage."""
    response = mock.Mock()
    response.status_code = status
    response.headers = headers or {}
    response.encoding = encoding
    response.is_redirect = is_redirect
    response.close = mock.Mock()
    response.iter_content = mock.Mock(
        return_value=iter(chunks if chunks is not None else [body])
    )
    return response


# ---------------------------------------------------------------------------
# Parsers
# ---------------------------------------------------------------------------


class ParserTests(SimpleTestCase):
    def test_unsupported_extension_rejected(self):
        upload = SimpleUploadedFile("resume.exe", b"binary stuff")
        with self.assertRaises(parsers.ResumeParseError) as ctx:
            parsers.extract_text_from_file(upload)
        self.assertIn("Unsupported file type", str(ctx.exception))

    def test_oversize_file_rejected(self):
        fake = SimpleNamespace(name="big.txt", size=parsers.MAX_FILE_SIZE + 1)
        with self.assertRaises(parsers.ResumeParseError) as ctx:
            parsers.validate_upload(fake)
        self.assertIn("too large", str(ctx.exception))

    def test_plain_text_extracted_and_cleaned(self):
        upload = SimpleUploadedFile(
            "resume.txt", b"John   Doe\n\n\n\nSoftware\tEngineer\n" + b"x " * 100
        )
        text = parsers.extract_text_from_file(upload)
        self.assertTrue(text.startswith("John Doe\n\nSoftware Engineer"))

    def test_too_little_text_rejected(self):
        upload = SimpleUploadedFile("tiny.txt", b"Hi there")
        with self.assertRaises(parsers.ResumeParseError) as ctx:
            parsers.extract_text_from_file(upload)
        self.assertIn("couldn't read enough text", str(ctx.exception))

    def test_corrupted_pdf_rejected(self):
        upload = SimpleUploadedFile("broken.pdf", b"%PDF-1.4 no text here")
        with self.assertRaises(parsers.ResumeParseError) as ctx:
            parsers.extract_text_from_file(upload)
        self.assertIn("could not be read", str(ctx.exception))

    def test_pdf_text_extracted(self):
        pdf_bytes = build_text_pdf("Hello PDF resume. " * 10)
        upload = SimpleUploadedFile("resume.pdf", pdf_bytes)
        text = parsers.extract_text_from_file(upload)
        self.assertIn("Hello PDF resume", text)

    def test_docx_paragraphs_and_merged_table_cells(self):
        document = Document()
        document.add_paragraph("Some resume content. " * 10)

        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).merge(table.cell(0, 1)).text = "HEADER"
        table.cell(1, 0).text = "Languages"
        table.cell(1, 1).text = "Python"

        buffer = BytesIO()
        document.save(buffer)
        upload = SimpleUploadedFile("resume.docx", buffer.getvalue())

        text = parsers.extract_text_from_file(upload)
        self.assertIn("Some resume content", text)
        self.assertIn("Languages", text)
        # The merged cell must appear exactly once, not once per grid column.
        self.assertEqual(text.count("HEADER"), 1)


# ---------------------------------------------------------------------------
# Scraper: URL validation (the SSRF guard)
# ---------------------------------------------------------------------------


class ScraperUrlValidationTests(SimpleTestCase):
    def test_rejects_private_and_local_addresses(self):
        bad_urls = [
            "http://127.0.0.1:8000/",
            "http://localhost:8000/",
            "http://169.254.169.254/latest/meta-data/",
            "http://192.168.1.1/",
            "http://10.0.0.5/",
            "http://[::1]/",
        ]
        for url in bad_urls:
            with self.subTest(url=url):
                with self.assertRaises(JobFetchError):
                    validate_public_url(url)

    def test_rejects_non_http_schemes(self):
        for url in ["file:///etc/passwd", "ftp://example.com/x", "gopher://example.com"]:
            with self.subTest(url=url):
                with self.assertRaises(JobFetchError):
                    validate_public_url(url)

    def test_rejects_url_without_host(self):
        with self.assertRaises(JobFetchError):
            validate_public_url("not a url")

    def test_accepts_public_ip(self):
        # Literal public IP so no external DNS is needed in tests.
        self.assertEqual(
            validate_public_url("http://93.184.216.34/jobs"),
            "http://93.184.216.34/jobs",
        )


# ---------------------------------------------------------------------------
# Scraper: fetching and extraction
# ---------------------------------------------------------------------------


class ScraperFetchTests(SimpleTestCase):
    def test_redirect_to_internal_address_blocked(self):
        redirect = make_http_response(
            status=302,
            headers={"Location": "http://127.0.0.1/secret"},
            is_redirect=True,
        )
        with mock.patch("resumes.scraper.requests.get", return_value=redirect):
            with self.assertRaises(JobFetchError) as ctx:
                fetch_job_page("http://93.184.216.34/job")
        self.assertIn("private or local", str(ctx.exception))

    def test_too_many_redirects(self):
        redirect = make_http_response(
            status=302,
            headers={"Location": "http://93.184.216.34/next"},
            is_redirect=True,
        )
        with mock.patch("resumes.scraper.MAX_REDIRECTS", 1), mock.patch(
            "resumes.scraper.requests.get", return_value=redirect
        ):
            with self.assertRaises(JobFetchError) as ctx:
                fetch_job_page("http://93.184.216.34/job")
        self.assertIn("Too many redirects", str(ctx.exception))

    def test_404_friendly_error(self):
        response = make_http_response(status=404)
        with mock.patch("resumes.scraper.requests.get", return_value=response):
            with self.assertRaises(JobFetchError) as ctx:
                fetch_job_page("http://93.184.216.34/job")
        self.assertIn("not found", str(ctx.exception))

    def test_blocked_403_friendly_error(self):
        response = make_http_response(status=403)
        with mock.patch("resumes.scraper.requests.get", return_value=response):
            with self.assertRaises(JobFetchError) as ctx:
                fetch_job_page("http://93.184.216.34/job")
        self.assertIn("blocked automated access", str(ctx.exception))

    def test_non_html_content_type_rejected(self):
        response = make_http_response(
            status=200, headers={"Content-Type": "application/pdf"}
        )
        with mock.patch("resumes.scraper.requests.get", return_value=response):
            with self.assertRaises(JobFetchError) as ctx:
                fetch_job_page("http://93.184.216.34/job.pdf")
        self.assertIn("not point to a readable web page", str(ctx.exception))

    def test_oversized_page_rejected(self):
        response = make_http_response(
            status=200,
            headers={"Content-Type": "text/html"},
            chunks=[b"x" * 200],
        )
        with mock.patch("resumes.scraper.MAX_HTML_BYTES", 100), mock.patch(
            "resumes.scraper.requests.get", return_value=response
        ):
            with self.assertRaises(JobFetchError) as ctx:
                fetch_job_page("http://93.184.216.34/job")
        self.assertIn("too large", str(ctx.exception))

    def test_timeout_mapped(self):
        with mock.patch("resumes.scraper.requests.get", side_effect=requests.Timeout):
            with self.assertRaises(JobFetchError) as ctx:
                fetch_job_page("http://93.184.216.34/job")
        self.assertIn("took too long", str(ctx.exception))

    def test_connection_error_mapped(self):
        with mock.patch(
            "resumes.scraper.requests.get", side_effect=requests.ConnectionError
        ):
            with self.assertRaises(JobFetchError) as ctx:
                fetch_job_page("http://93.184.216.34/job")
        self.assertIn("Could not reach", str(ctx.exception))

    def test_happy_path_extracts_description(self):
        html = (
            "<html><body><article>"
            "<p>" + "We are hiring a Python engineer to build Django services. " * 6 + "</p>"
            "<p>" + "Requirements: PostgreSQL, Docker, AWS and strong testing skills. " * 6 + "</p>"
            "</article></body></html>"
        ).encode()
        response = make_http_response(
            status=200, body=html, headers={"Content-Type": "text/html"}
        )
        with mock.patch("resumes.scraper.requests.get", return_value=response):
            text = fetch_job_description("http://93.184.216.34/job")
        self.assertIn("Python engineer", text)

    def test_empty_extraction_friendly_error(self):
        with self.assertRaises(JobFetchError) as ctx:
            extract_job_description("<html><body></body></html>")
        self.assertIn("Couldn't find a readable job description", str(ctx.exception))

    def test_short_extraction_friendly_error(self):
        with mock.patch("resumes.scraper.trafilatura.extract", return_value="too short"):
            with self.assertRaises(JobFetchError) as ctx:
                extract_job_description("<html><body><p>x</p></body></html>")
        self.assertIn("too short to analyze", str(ctx.exception))

    def test_trafilatura_crash_mapped(self):
        with mock.patch(
            "resumes.scraper.trafilatura.extract", side_effect=RuntimeError("boom")
        ):
            with self.assertRaises(JobFetchError):
                extract_job_description("<html><body><p>x</p></body></html>")


# ---------------------------------------------------------------------------
# Similarity (embeddings + cosine)
# ---------------------------------------------------------------------------


class SimilarityTests(SimpleTestCase):
    def test_cosine_identical_vectors(self):
        self.assertAlmostEqual(_cosine_similarity([1.0, 0.0], [1.0, 0.0]), 1.0)

    def test_cosine_orthogonal_vectors(self):
        self.assertAlmostEqual(_cosine_similarity([1.0, 0.0], [0.0, 1.0]), 0.0)

    def test_cosine_degenerate_vector(self):
        self.assertIsNone(_cosine_similarity([0.0, 0.0], [1.0, 0.0]))

    def test_empty_inputs_return_none(self):
        self.assertIsNone(semantic_similarity("", "job text"))
        self.assertIsNone(semantic_similarity("resume text", ""))

    def test_computes_cosine_from_two_embeddings(self):
        client = mock.Mock()
        client.models.embed_content.return_value = mock.Mock(
            embeddings=[mock.Mock(values=[1.0, 0.0]), mock.Mock(values=[1.0, 0.0])]
        )

        with mock.patch("resumes.similarity.get_client", return_value=client):
            similarity = semantic_similarity(RESUME_TEXT, "job text " * 50)

        self.assertAlmostEqual(similarity, 1.0)

        # Two separate Content objects must be sent, otherwise
        # gemini-embedding-2 returns one aggregated vector.
        contents = client.models.embed_content.call_args.kwargs["contents"]
        self.assertEqual(len(contents), 2)

    def test_embedding_failure_returns_none(self):
        with mock.patch(
            "resumes.similarity.get_client", side_effect=RuntimeError("no client")
        ):
            self.assertIsNone(semantic_similarity(RESUME_TEXT, "job text " * 50))


# ---------------------------------------------------------------------------
# Analyzer
# ---------------------------------------------------------------------------


class AnalyzerTests(SimpleTestCase):
    def test_general_mode_and_normalization(self):
        client = make_mock_client(fake_result_json(72))

        with mock.patch("resumes.analyzer.get_client", return_value=client), mock.patch(
            "resumes.analyzer.semantic_similarity", return_value=None
        ):
            result = analyze_resume(RESUME_TEXT)

        self.assertEqual(result["mode"], "GENERAL")
        self.assertEqual(result["overall_score"], 72)
        self.assertEqual(result["ai_model"], settings.GEMINI_MODEL)
        self.assertEqual(result["prompt_tokens"], 100)
        self.assertEqual(result["output_tokens"], 50)
        self.assertIsNone(result["semantic_similarity"])

        # Keywords are lowercased, stripped and deduped.
        self.assertEqual(
            result["result"]["matched_keywords"], ["python", "django"]
        )
        self.assertEqual(result["result"]["missing_keywords"], ["kubernetes"])

    def test_job_match_mode_propagates_similarity(self):
        client = make_mock_client(fake_result_json(40))

        with mock.patch("resumes.analyzer.get_client", return_value=client), mock.patch(
            "resumes.analyzer.semantic_similarity", return_value=0.73
        ):
            result = analyze_resume(RESUME_TEXT, "Senior Python engineer wanted. " * 20)

        self.assertEqual(result["mode"], "JOB_MATCH")
        self.assertAlmostEqual(result["semantic_similarity"], 0.73)

    def test_request_shape(self):
        client = make_mock_client(fake_result_json())

        with mock.patch("resumes.analyzer.get_client", return_value=client), mock.patch(
            "resumes.analyzer.semantic_similarity", return_value=None
        ):
            analyze_resume(RESUME_TEXT)

        kwargs = client.interactions.create.call_args.kwargs
        self.assertEqual(kwargs["response_format"]["mime_type"], "application/json")
        self.assertIn("overall_score", kwargs["response_format"]["schema"]["properties"])
        self.assertEqual(kwargs["generation_config"]["thinking_level"], "low")

    def test_invalid_json_raises(self):
        client = make_mock_client("this is not json")

        with mock.patch("resumes.analyzer.get_client", return_value=client), mock.patch(
            "resumes.analyzer.semantic_similarity", return_value=None
        ):
            with self.assertRaises(AIAnalysisError) as ctx:
                analyze_resume(RESUME_TEXT)
        self.assertIn("unexpected result", str(ctx.exception))

    def test_empty_output_raises(self):
        client = make_mock_client("")

        with mock.patch("resumes.analyzer.get_client", return_value=client), mock.patch(
            "resumes.analyzer.semantic_similarity", return_value=None
        ):
            with self.assertRaises(AIAnalysisError) as ctx:
                analyze_resume(RESUME_TEXT)
        self.assertIn("empty response", str(ctx.exception))

    def test_rate_limit_mapped(self):
        client = mock.Mock()
        client.interactions.create.side_effect = errors.ClientError(
            429, {"error": {"message": "quota"}}
        )

        with mock.patch("resumes.analyzer.get_client", return_value=client), mock.patch(
            "resumes.analyzer.semantic_similarity", return_value=None
        ):
            with self.assertRaises(AIAnalysisError) as ctx:
                analyze_resume(RESUME_TEXT)
        self.assertIn("limit", str(ctx.exception))

    def test_server_error_mapped(self):
        client = mock.Mock()
        client.interactions.create.side_effect = errors.ServerError(
            500, {"error": {"message": "boom"}}
        )

        with mock.patch("resumes.analyzer.get_client", return_value=client), mock.patch(
            "resumes.analyzer.semantic_similarity", return_value=None
        ):
            with self.assertRaises(AIAnalysisError) as ctx:
                analyze_resume(RESUME_TEXT)
        self.assertIn("temporarily unavailable", str(ctx.exception))

    @override_settings(GEMINI_API_KEY=None)
    def test_missing_api_key_raises_config_error(self):
        with self.assertRaises(AIAnalysisError) as ctx:
            analyze_resume(RESUME_TEXT)
        self.assertIn("not configured", str(ctx.exception))

    def test_truncate_marks_cut_text(self):
        self.assertEqual(_truncate("short", 100), "short")
        truncated = _truncate("x" * 200, 50)
        self.assertEqual(len(truncated), 50)
        self.assertTrue(truncated.endswith("…"))

    def test_build_prompt_includes_inputs_and_similarity(self):
        prompt = build_prompt("RESUME-MARKER", "JOB-MARKER", similarity=0.73)
        self.assertIn("RESUME-MARKER", prompt)
        self.assertIn("JOB-MARKER", prompt)
        self.assertIn("0.73", prompt)
        self.assertIn("MODE: job match", prompt)

    def test_build_prompt_general_mode(self):
        prompt = build_prompt("RESUME-MARKER")
        self.assertIn("MODE: general analysis", prompt)
        self.assertNotIn("job_description", prompt)


# ---------------------------------------------------------------------------
# Forms
# ---------------------------------------------------------------------------


class FormTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.user = User.objects.create_user(email="form_a@example.com", password="x")
        cls.other = User.objects.create_user(email="form_b@example.com", password="x")
        cls.resume = Resume.objects.create(
            user=cls.user, text=RESUME_TEXT, original_filename="a.txt"
        )
        Resume.objects.create(
            user=cls.other, text=RESUME_TEXT, original_filename="b.txt"
        )

    def test_no_source_invalid(self):
        form = ResumeAnalysisForm(self.user, data={})
        self.assertFalse(form.is_valid())
        self.assertIn("Provide a resume", str(form.non_field_errors()))

    def test_multiple_sources_invalid(self):
        form = ResumeAnalysisForm(
            self.user,
            data={"resume_text": RESUME_TEXT},
            files={"resume_file": SimpleUploadedFile("r.txt", b"data")},
        )
        self.assertFalse(form.is_valid())
        self.assertIn("Choose only one resume source", str(form.non_field_errors()))

    def test_existing_resume_valid(self):
        form = ResumeAnalysisForm(self.user, data={"resume_choice": self.resume.pk})
        self.assertTrue(form.is_valid(), form.errors)

    def test_pasted_text_valid(self):
        form = ResumeAnalysisForm(self.user, data={"resume_text": RESUME_TEXT})
        self.assertTrue(form.is_valid(), form.errors)

    def test_short_text_invalid(self):
        form = ResumeAnalysisForm(self.user, data={"resume_text": "too short"})
        self.assertFalse(form.is_valid())
        self.assertIn("resume_text", form.errors)

    def test_wrong_extension_invalid(self):
        form = ResumeAnalysisForm(
            self.user,
            data={},
            files={"resume_file": SimpleUploadedFile("resume.exe", b"binary")},
        )
        self.assertFalse(form.is_valid())
        self.assertIn("resume_file", form.errors)

    def test_querysets_scoped_to_user(self):
        form = ResumeAnalysisForm(self.user)
        self.assertEqual(form.fields["resume_choice"].queryset.count(), 1)
        self.assertEqual(form.fields["job_application"].queryset.count(), 0)


# ---------------------------------------------------------------------------
# Analysis flow (views, with AI and scraper mocked)
# ---------------------------------------------------------------------------


@override_settings(MEDIA_ROOT=TEMP_MEDIA)
class AnalyzeFlowTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.user = User.objects.create_user(email="flow@example.com", password="x")
        cls.application = JobApplication.objects.create(
            user=cls.user,
            company="Acme",
            job_title="Backend Engineer",
            application_date="2026-01-01",
            job_url="https://93.184.216.34/job",
        )

    def setUp(self):
        self.client.force_login(self.user)

    def test_login_required(self):
        self.client.logout()
        response = self.client.get(reverse("resumes:list"))
        self.assertEqual(response.status_code, 302)
        self.assertIn("/accounts/login/", response.url)

    def test_analyze_page_renders(self):
        response = self.client.get(reverse("resumes:analyze"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Analyze Your Resume")
        self.assertContains(response, 'enctype="multipart/form-data"')

    def test_success_creates_analysis_with_mapping(self):
        with mock.patch("resumes.views.analyze_resume", side_effect=fake_analyze) as fake:
            response = self.client.post(
                reverse("resumes:analyze"), {"resume_text": RESUME_TEXT}
            )

        self.assertEqual(response.status_code, 302)
        analysis = ResumeAnalysis.objects.get()
        self.assertEqual(analysis.mode, "GENERAL")
        self.assertEqual(analysis.overall_score, 55)
        self.assertEqual(analysis.ai_model, "mock-model")
        self.assertEqual(analysis.prompt_tokens, 10)
        self.assertEqual(analysis.output_tokens, 20)
        self.assertEqual(analysis.resume.original_filename, "Pasted text")
        # The analyzer receives the full text, no job description.
        self.assertIsNone(fake.call_args.args[1])

    def test_paste_wins_over_link(self):
        with mock.patch("resumes.views.analyze_resume", side_effect=fake_analyze), mock.patch(
            "resumes.views.fetch_job_description"
        ) as fake_fetch:
            response = self.client.post(
                reverse("resumes:analyze"),
                {
                    "resume_text": RESUME_TEXT,
                    "job_url": "https://93.184.216.34/job",
                    "job_description": "We need a Python engineer. " * 10,
                },
            )

        self.assertEqual(response.status_code, 302)
        self.assertFalse(fake_fetch.called)
        analysis = ResumeAnalysis.objects.get()
        self.assertEqual(analysis.mode, "JOB_MATCH")
        self.assertEqual(analysis.job_url, "https://93.184.216.34/job")

    def test_fetch_failure_shows_field_error(self):
        with mock.patch(
            "resumes.views.fetch_job_description",
            side_effect=JobFetchError("This website blocked automated access (HTTP 403)."),
        ), mock.patch("resumes.views.analyze_resume", side_effect=fake_analyze):
            response = self.client.post(
                reverse("resumes:analyze"),
                {"resume_text": RESUME_TEXT, "job_url": "https://93.184.216.34/job"},
            )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "blocked automated access")
        self.assertEqual(ResumeAnalysis.objects.count(), 0)
        # No resume row either: nothing was persisted before the fetch failed.
        self.assertEqual(Resume.objects.count(), 0)

    def test_ai_failure_keeps_resume(self):
        with mock.patch(
            "resumes.views.analyze_resume",
            side_effect=AIAnalysisError("The AI service is at its limit right now."),
        ):
            response = self.client.post(
                reverse("resumes:analyze"), {"resume_text": RESUME_TEXT}
            )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "limit right now")
        self.assertEqual(ResumeAnalysis.objects.count(), 0)
        self.assertEqual(Resume.objects.count(), 1)

    @override_settings(ATS_DAILY_LIMIT=0)
    def test_daily_limit_blocks_before_any_work(self):
        with mock.patch("resumes.views.analyze_resume", side_effect=fake_analyze) as fake:
            response = self.client.post(
                reverse("resumes:analyze"), {"resume_text": RESUME_TEXT}
            )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Daily limit reached")
        self.assertFalse(fake.called)
        self.assertEqual(Resume.objects.count(), 0)

    def test_resume_preselected_from_query(self):
        resume = Resume.objects.create(
            user=self.user, text=RESUME_TEXT, original_filename="saved.txt"
        )
        response = self.client.get(reverse("resumes:analyze") + f"?resume={resume.pk}")
        self.assertContains(response, f'value="{resume.pk}" selected')

    def test_application_link_used_when_no_paste(self):
        with mock.patch("resumes.views.analyze_resume", side_effect=fake_analyze), mock.patch(
            "resumes.views.fetch_job_description", return_value="Fetched JD text. " * 20
        ) as fake_fetch:
            response = self.client.post(
                reverse("resumes:analyze"),
                {"resume_text": RESUME_TEXT, "job_application": self.application.pk},
            )

        self.assertEqual(response.status_code, 302)
        fake_fetch.assert_called_once_with(self.application.job_url)
        analysis = ResumeAnalysis.objects.get()
        self.assertEqual(analysis.job_application, self.application)

    def test_upload_creates_resume_with_file(self):
        upload = SimpleUploadedFile(
            "resume.txt", ("Jane Doe resume text. " * 20).encode()
        )
        with mock.patch("resumes.views.analyze_resume", side_effect=fake_analyze):
            response = self.client.post(
                reverse("resumes:analyze"), {"resume_file": upload}
            )

        self.assertEqual(response.status_code, 302)
        resume = Resume.objects.get()
        self.assertTrue(resume.file)
        self.assertEqual(resume.original_filename, "resume.txt")


# ---------------------------------------------------------------------------
# Access control + dashboard
# ---------------------------------------------------------------------------


class AccessControlTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.owner = User.objects.create_user(email="owner@example.com", password="x")
        cls.intruder = User.objects.create_user(email="intruder@example.com", password="x")

        cls.resume = Resume.objects.create(
            user=cls.owner, text=RESUME_TEXT, original_filename="owner.txt"
        )
        cls.analysis = ResumeAnalysis.objects.create(
            resume=cls.resume,
            mode=ResumeAnalysis.Mode.GENERAL,
            overall_score=70,
            result=json.loads(fake_result_json(70)),
            ai_model="mock-model",
        )

    def test_other_user_cannot_view_detail(self):
        self.client.force_login(self.intruder)
        response = self.client.get(
            reverse("resumes:analysis_detail", args=[self.analysis.pk])
        )
        self.assertEqual(response.status_code, 404)

    def test_other_user_cannot_delete(self):
        self.client.force_login(self.intruder)
        response = self.client.post(
            reverse("resumes:analysis_delete", args=[self.analysis.pk])
        )
        self.assertEqual(response.status_code, 404)
        self.assertTrue(ResumeAnalysis.objects.filter(pk=self.analysis.pk).exists())

    def test_other_user_cannot_download(self):
        self.client.force_login(self.intruder)
        response = self.client.get(
            reverse("resumes:resume_download", args=[self.resume.pk])
        )
        self.assertEqual(response.status_code, 404)

    def test_list_shows_only_own_rows(self):
        other_resume = Resume.objects.create(
            user=self.intruder, text=RESUME_TEXT, original_filename="intruder.txt"
        )
        ResumeAnalysis.objects.create(
            resume=other_resume,
            mode=ResumeAnalysis.Mode.GENERAL,
            overall_score=10,
            result={},
            ai_model="mock-model",
        )

        self.client.force_login(self.owner)
        response = self.client.get(reverse("resumes:list"))
        self.assertContains(response, "owner.txt")
        self.assertNotContains(response, "intruder.txt")


class DashboardTests(TestCase):
    def test_dashboard_without_analysis_shows_cta(self):
        user = User.objects.create_user(email="dash_empty@example.com", password="x")
        self.client.force_login(user)
        response = self.client.get(reverse("dashboard:dashboard"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Score your resume")

    def test_dashboard_with_analysis_shows_latest_score(self):
        user = User.objects.create_user(email="dash_full@example.com", password="x")
        resume = Resume.objects.create(
            user=user, text=RESUME_TEXT, original_filename="dash.txt"
        )
        ResumeAnalysis.objects.create(
            resume=resume,
            mode=ResumeAnalysis.Mode.JOB_MATCH,
            overall_score=82,
            result={},
            ai_model="mock-model",
        )

        self.client.force_login(user)
        response = self.client.get(reverse("dashboard:dashboard"))
        self.assertContains(response, "Latest score")
        self.assertContains(response, "82/100")


# ---------------------------------------------------------------------------
# Files: download, delete and disk cleanup
# ---------------------------------------------------------------------------


@override_settings(MEDIA_ROOT=TEMP_MEDIA)
class FileLifecycleTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(email="files@example.com", password="x")
        self.client.force_login(self.user)

    def _upload_resume(self):
        upload = SimpleUploadedFile("resume.txt", ("File resume. " * 20).encode())
        return Resume.objects.create(
            user=self.user,
            file=upload,
            text="File resume.",
            original_filename="resume.txt",
        )

    def test_download_uploaded_file(self):
        resume = self._upload_resume()
        response = self.client.get(reverse("resumes:resume_download", args=[resume.pk]))
        self.assertEqual(response.status_code, 200)
        self.assertIn("attachment", response["Content-Disposition"])

    def test_download_pasted_resume_redirects(self):
        resume = Resume.objects.create(
            user=self.user, text=RESUME_TEXT, original_filename="Pasted text"
        )
        response = self.client.get(reverse("resumes:resume_download", args=[resume.pk]))
        self.assertEqual(response.status_code, 302)

    def test_resume_delete_removes_file_from_disk(self):
        resume = self._upload_resume()
        path = resume.file.path
        self.assertTrue(os.path.exists(path))

        response = self.client.post(reverse("resumes:resume_delete", args=[resume.pk]))

        self.assertEqual(response.status_code, 302)
        self.assertFalse(Resume.objects.filter(pk=resume.pk).exists())
        self.assertFalse(os.path.exists(path))

    def test_resume_delete_cascades_analyses(self):
        resume = self._upload_resume()
        ResumeAnalysis.objects.create(
            resume=resume,
            mode=ResumeAnalysis.Mode.GENERAL,
            overall_score=60,
            result={},
            ai_model="mock-model",
        )

        self.client.post(reverse("resumes:resume_delete", args=[resume.pk]))

        self.assertEqual(ResumeAnalysis.objects.count(), 0)

    def test_analysis_delete(self):
        resume = Resume.objects.create(
            user=self.user, text=RESUME_TEXT, original_filename="Pasted text"
        )
        analysis = ResumeAnalysis.objects.create(
            resume=resume,
            mode=ResumeAnalysis.Mode.GENERAL,
            overall_score=60,
            result={},
            ai_model="mock-model",
        )

        response = self.client.post(
            reverse("resumes:analysis_delete", args=[analysis.pk])
        )

        self.assertEqual(response.status_code, 302)
        self.assertEqual(ResumeAnalysis.objects.count(), 0)
        # The resume itself survives.
        self.assertTrue(Resume.objects.filter(pk=resume.pk).exists())
