"""
Resume text extraction.

Turns uploaded resume files (PDF, DOCX, TXT) into clean plain text that can
be sent to the AI.

This module deliberately has no Django imports: it accepts any file-like
object with a `.name` attribute, so it can be unit tested in isolation and
reused outside of a request cycle (management commands, scripts, tests).
"""

from pathlib import Path

from docx import Document
from pypdf import PdfReader

from .text_utils import clean_text

ALLOWED_EXTENSIONS = {".pdf", ".docx", ".txt"}

# 5 MB is far beyond any real resume; anything bigger is a mistake or abuse.
MAX_FILE_SIZE = 5 * 1024 * 1024

# Below this, the file is almost certainly a scanned image without a text
# layer. We cannot parse those (yet), so we ask the user to paste text.
MIN_TEXT_LENGTH = 100


class ResumeParseError(Exception):
    """User-facing error raised when a resume cannot be turned into text."""


def validate_upload(uploaded_file):
    """
    Check the file extension and size before parsing.

    Returns the lowercase extension, e.g. ".pdf".
    Raises ResumeParseError with a message safe to show to the user.
    """
    extension = Path(uploaded_file.name).suffix.lower()

    if extension not in ALLOWED_EXTENSIONS:
        allowed = ", ".join(sorted(ALLOWED_EXTENSIONS))
        raise ResumeParseError(
            f"Unsupported file type '{extension or 'unknown'}'. "
            f"Please upload one of: {allowed}."
        )

    size = getattr(uploaded_file, "size", None)
    if size is not None and size > MAX_FILE_SIZE:
        raise ResumeParseError(
            f"This file is too large ({size / 1024 / 1024:.1f} MB). "
            "The maximum allowed size is 5 MB."
        )

    return extension


def extract_pdf_text(uploaded_file):
    """Extract text from a PDF using pypdf, page by page."""
    try:
        reader = PdfReader(uploaded_file)
    except Exception as exc:
        raise ResumeParseError(
            "This PDF could not be read. It may be corrupted — try "
            "re-exporting it or paste your resume text instead."
        ) from exc

    if reader.is_encrypted:
        # Many PDFs are "encrypted" with an empty user password; try that.
        try:
            if not reader.decrypt(""):
                raise ResumeParseError(
                    "This PDF is password protected. Please remove the "
                    "password or paste your resume text instead."
                )
        except ResumeParseError:
            raise
        except Exception as exc:
            raise ResumeParseError(
                "This PDF is password protected. Please remove the "
                "password or paste your resume text instead."
            ) from exc

    pages = []
    for page in reader.pages:
        # extract_text() returns None on pages without a text layer.
        pages.append(page.extract_text() or "")

    return "\n".join(pages)


def extract_docx_text(uploaded_file):
    """Extract text from a DOCX, including tables (common resume layout)."""
    try:
        document = Document(uploaded_file)
    except Exception as exc:
        raise ResumeParseError(
            "This DOCX file could not be read. It may be corrupted — "
            "try re-exporting it or paste your resume text instead."
        ) from exc

    parts = [paragraph.text for paragraph in document.paragraphs]

    # Merged table cells appear multiple times in row.cells (once per grid
    # column they span). We hold references to the underlying <w:tc> elements
    # and compare with `is`: using id() is unsafe because a temporary proxy
    # can be garbage-collected and its address reused by a later cell.
    seen_cells = []
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if any(cell._tc is seen for seen in seen_cells):
                    continue
                seen_cells.append(cell._tc)
                parts.append(cell.text)

    return "\n".join(parts)


def extract_text_from_file(uploaded_file):
    """
    Validate, dispatch by extension, clean, and sanity-check an upload.

    Accepts a Django UploadedFile or any file-like object with `.name`.
    Returns cleaned text or raises ResumeParseError with a user-facing message.
    """
    extension = validate_upload(uploaded_file)

    # File pointers may have moved during validation; always start at zero.
    if hasattr(uploaded_file, "seek"):
        uploaded_file.seek(0)

    if extension == ".pdf":
        text = extract_pdf_text(uploaded_file)
    elif extension == ".docx":
        text = extract_docx_text(uploaded_file)
    else:  # .txt
        raw = uploaded_file.read()
        text = raw.decode("utf-8", errors="replace") if isinstance(raw, bytes) else raw

    cleaned = clean_text(text)

    if len(cleaned) < MIN_TEXT_LENGTH:
        raise ResumeParseError(
            "We couldn't read enough text from this file — it may be a "
            "scanned image or empty. Please paste your resume text instead."
        )

    return cleaned
